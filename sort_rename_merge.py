#!/usr/bin/env python3
"""
sort_rename_merge.py

Audit-ready tool to read multiple Excel files, match PDFs, sort by Beneficial Owner (BO),
rename, archive, and optionally merge PDFs.

Requirements:
- Python 3.9+
- Libraries: openpyxl, PyPDF2
- No runtime user interaction
- Works on Windows and Linux
"""

from __future__ import annotations

import json
import re
import shutil
import sys
import time
import traceback
import unicodedata
from collections import defaultdict
from copy import deepcopy
from datetime import datetime, date
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from PyPDF2 import PdfMerger, PdfReader, PdfWriter


# =========================
# Global Log Buffer
# =========================

LOG_BUFFER: List[str] = []


def log(level: str, message: str) -> None:
    """Append a log message to the in-memory log buffer."""
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    LOG_BUFFER.append(f"{timestamp} [{level}] {message}")


# =========================
# Default Settings
# =========================

DEFAULT_SETTINGS: Dict[str, Any] = {
    "columns": {
        "isin": "Financial Instrument",
        "date": "Paymt. Date",
        "bo": "BO Name",
        "pdf_base": "Clearstream Request ID",
    },
    "actions": {
        "create_outputs": True,
        "sort_into_bo_folders": True,
        "rename_pdfs": True,
        "merge_all": True,
        "per_bo_merge": False,
        "cleanup_input": False,
    },
    "lookup": {
        "strategies": ["exact"],
        "case_insensitive": True,
    },
    "sorting": {
        "mode": "default",
        "bo_lookup_file": "input/bo_lookup.xlsx",
    },
    "runtime": {
        "mode": "normal",
        "normal": {
            "retries": 1,
            "delay_ms": 0,
        },
        "robust": {
            "retries": 5,
            "delay_ms": 250,
        },
    },
    "naming": {
        "date_format": "%Y-%m-%d",
        "output_filename_pattern": "{isin}_{date}.pdf",
        "bo_folder_pattern": "{bo}",
        "merged_all_filename": "merged_all.pdf",
        "merged_bo_filename_pattern": "merged_{bo}.pdf",
    },
    "paths": {
        "excel_input": "input/excel",
        "pdf_input": "input/pdf_inbox",
        "archive": "archive",
        "snapshot_excel": "input_snapshot/excel",
        "snapshot_pdf": "input_snapshot/pdf_inbox",
        "individual_output": "sorted_by_bo",
        "merged_output": "merged",
    },
    "sanitize_filenames": True,
}


# =========================
# Settings Loader
# =========================

def strip_jsonc_comments(text: str) -> str:
    """Remove // and /* */ comments from JSONC text."""
    text = re.sub(r"//.*?$", "", text, flags=re.MULTILINE)
    text = re.sub(r"/\*.*?\*/", "", text, flags=re.DOTALL)
    return text


def load_settings(base_dir: Path) -> Dict[str, Any]:
    """Load settings.json or settings.jsonc, otherwise return defaults."""
    json_file = base_dir / "settings.json"
    jsonc_file = base_dir / "settings.jsonc"

    if json_file.exists():
        try:
            with json_file.open("r", encoding="utf-8") as f:
                data = json.load(f)
            normalize_legacy_settings(data)
            log("INFO", "Loaded settings.json")
            return deep_merge(DEFAULT_SETTINGS, data)
        except Exception as e:
            log("ERROR", f"Failed to load settings.json: {e}")

    if jsonc_file.exists():
        try:
            with jsonc_file.open("r", encoding="utf-8") as f:
                raw = f.read()
            stripped = strip_jsonc_comments(raw)
            data = json.loads(stripped)
            normalize_legacy_settings(data)
            log("INFO", "Loaded settings.jsonc")
            return deep_merge(DEFAULT_SETTINGS, data)
        except Exception as e:
            log("ERROR", f"Failed to load settings.jsonc: {e}")

    log("INFO", "No settings file found. Using default settings.")
    return deepcopy(DEFAULT_SETTINGS)


def deep_merge(base: Dict[str, Any], override: Dict[str, Any]) -> Dict[str, Any]:
    """Deep merge two dictionaries."""
    result = dict(base)
    for k, v in override.items():
        if isinstance(v, dict) and k in result and isinstance(result[k], dict):
            result[k] = deep_merge(result[k], v)
        else:
            result[k] = v
    return result


def normalize_legacy_settings(data: Dict[str, Any]) -> None:
    """Translate legacy settings keys to the current schema."""
    actions = data.get("actions")
    if not isinstance(actions, dict):
        return

    if "merge" in actions and "merge_all" not in actions:
        actions["merge_all"] = actions["merge"]

    actions.pop("move_pdfs", None)


# =========================
# Utility Functions
# =========================

INVALID_CHARS = r'[<>:"/\\|?*]'


def sanitize_filename(name: str) -> str:
    """Sanitize filename for Windows compatibility."""
    name = unicodedata.normalize("NFKD", name)
    name = re.sub(INVALID_CHARS, "_", name)
    name = name.strip()
    return name


def normalize_lookup_key(value: str) -> str:
    """Normalize BO names for lookup matching."""
    normalized = unicodedata.normalize("NFKD", str(value))
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return normalized.strip().lower()


def render_template(template: str, values: Dict[str, str], sanitize: bool) -> str:
    """Render a configurable filename/folder template."""
    rendered = template.format(**values)
    return sanitize_filename(rendered) if sanitize else rendered.strip()


def get_runtime_options(settings: Dict[str, Any]) -> Dict[str, Any]:
    """Resolve runtime behavior from the selected mode."""
    runtime = settings["runtime"]
    mode = runtime.get("mode", "normal")
    selected = runtime.get(mode, runtime["normal"])
    return {
        "mode": mode,
        "retries": max(1, int(selected.get("retries", 1))),
        "delay_ms": max(0, int(selected.get("delay_ms", 0))),
    }


def run_with_retries(
    operation_name: str,
    func: Any,
    runtime_options: Dict[str, Any],
) -> Any:
    """Run a file operation with optional retries for slower/locked drives."""
    attempts = runtime_options["retries"]
    delay_seconds = runtime_options["delay_ms"] / 1000.0
    last_error: Optional[Exception] = None

    for attempt in range(1, attempts + 1):
        try:
            return func()
        except Exception as exc:
            last_error = exc
            if attempt >= attempts:
                break

            log(
                "WARNING",
                f"{operation_name} failed (attempt {attempt}/{attempts}): {exc}. Retrying...",
            )
            if delay_seconds > 0:
                time.sleep(delay_seconds)

    assert last_error is not None
    raise last_error


def determine_run_status(
    skipped_records: int,
    missing_pdfs: List[str],
    defective_pdfs: List[str],
    unprocessed_input_pdfs: List[str],
) -> str:
    """Determine the final run status for manifest and shell output."""
    if skipped_records == 0 and not missing_pdfs and not defective_pdfs and not unprocessed_input_pdfs:
        return "SUCCESS"
    return "SUCCESS WITH WARNINGS"


def print_run_summary(
    run_id: str,
    status: str,
    excel_count: int,
    pdf_count: int,
    total_rows: int,
    valid_records: int,
    skipped_records: int,
    missing_pdfs: List[str],
    defective_pdfs: List[str],
    unprocessed_input_pdfs: List[str],
    output_files: List[str],
    merge_outputs: List[str],
    run_dir: Path,
    duration_seconds: float,
) -> None:
    """Print a compact run summary to the shell."""
    print()
    print(f"Run finished: {run_id}")
    print(f"Status: {status}")
    print()
    print(f"Excel files: {excel_count}")
    print(f"PDF files: {pdf_count}")
    print(f"Rows processed: {total_rows}")
    print(f"Valid records: {valid_records}")
    print(f"Skipped records: {skipped_records}")
    print(f"Missing PDFs: {len(missing_pdfs)}")
    print(f"Defective PDFs: {len(defective_pdfs)}")
    print(f"Unprocessed inbox PDFs: {len(unprocessed_input_pdfs)}")
    print(f"Output files created: {len(output_files)}")
    print(f"Merge files created: {len(merge_outputs)}")
    print(f"Run folder: {run_dir}")
    print(f"Duration: {duration_seconds:.2f}s")


def ensure_unique_path(path: Path) -> Path:
    """Ensure filename uniqueness by adding _1, _2 suffixes."""
    if not path.exists():
        return path

    base = path.stem
    ext = path.suffix
    parent = path.parent

    counter = 1
    while True:
        new_path = parent / f"{base}_{counter}{ext}"
        if not new_path.exists():
            return new_path
        counter += 1


def parse_excel_date(value: Any) -> Optional[date]:
    """Parse Excel date value into a date."""
    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if isinstance(value, str):
        for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except Exception:
                continue

    return None


def build_lookup_bo_display_name(first_name: str, last_name: str) -> str:
    """Build a stable BO display name for folder and per-BO merge naming."""
    parts = [str(last_name).strip(), str(first_name).strip()]
    return "_".join(part for part in parts if part)


def load_bo_lastname_lookup(
    base_dir: Path,
    settings: Dict[str, Any],
    runtime_options: Dict[str, Any],
) -> Dict[str, Dict[str, str]]:
    """Load BO lookup data when lastname sorting is enabled."""
    sorting_cfg = settings.get("sorting", {})
    lookup_path = base_dir / Path(sorting_cfg.get("bo_lookup_file", "input/bo_lookup.xlsx"))

    if not lookup_path.exists():
        log("WARNING", f"BO lookup file not found: {lookup_path}. Falling back to default sorting.")
        return {}

    try:
        wb = run_with_retries(
            f"Open BO lookup workbook {lookup_path.name}",
            lambda path=lookup_path: load_workbook(path, read_only=True, data_only=True),
            runtime_options,
        )
    except Exception as e:
        log("WARNING", f"Failed to load BO lookup file {lookup_path}: {e}. Falling back to default sorting.")
        return {}

    lookup_map: Dict[str, Dict[str, str]] = {}

    try:
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            title = row[3] if len(row) > 3 else None
            first_name = row[4] if len(row) > 4 else None
            last_name = row[5] if len(row) > 5 else None

            if not title or not first_name or not last_name:
                continue

            normalized_bo_name = normalize_lookup_key(f"{title}{first_name}{last_name}")
            if not normalized_bo_name:
                continue

            first_name_str = str(first_name).strip()
            last_name_str = str(last_name).strip()
            lookup_map[normalized_bo_name] = {
                "sort_key": build_lookup_bo_display_name(first_name_str, last_name_str),
                "display_bo": build_lookup_bo_display_name(first_name_str, last_name_str),
            }
    except Exception as e:
        log("WARNING", f"Failed reading BO lookup rows from {lookup_path}: {e}. Falling back to default sorting.")
        lookup_map = {}
    finally:
        wb.close()

    return lookup_map


def resolve_bo_sorting_values(
    bo_name: str,
    bo_lookup_entry: Optional[Dict[str, str]],
) -> Tuple[str, str]:
    """Resolve sort key and BO display name for the current sorting mode."""
    if not bo_lookup_entry:
        return bo_name, bo_name

    sort_key = bo_lookup_entry.get("sort_key", "").strip() or bo_name
    display_bo = bo_lookup_entry.get("display_bo", "").strip() or bo_name
    return sort_key, display_bo


# =========================
# Excel Processing
# =========================

def read_excel_files(
    excel_files: List[Path], settings: Dict[str, Any], runtime_options: Dict[str, Any]
) -> Tuple[List[Dict[str, Any]], int]:
    """Read all Excel files and extract relevant rows."""
    columns = settings["columns"]
    required_columns = {
        columns["isin"],
        columns["date"],
        columns["bo"],
        columns["pdf_base"],
    }

    records: List[Dict[str, Any]] = []
    total_rows = 0

    for excel_file in excel_files:
        try:
            wb = run_with_retries(
                f"Open Excel workbook {excel_file.name}",
                lambda path=excel_file: load_workbook(path, read_only=True, data_only=True),
                runtime_options,
            )
            ws = wb.active

            header = [str(c.value).strip() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))[0:]]

            col_map = {name: idx for idx, name in enumerate(header)}
            missing_columns = sorted(col for col in required_columns if col not in col_map)

            if missing_columns:
                log(
                    "ERROR",
                    f"Skipping Excel {excel_file.name}: missing required columns: {', '.join(missing_columns)}",
                )
                wb.close()
                continue

            log("INFO", f"Processing Excel: {excel_file.name}")

            file_rows = 0
            valid_rows = 0

            for row in ws.iter_rows(min_row=2, values_only=True):
                total_rows += 1
                file_rows += 1

                if not row:
                    continue

                first = str(row[0]) if row[0] else ""
                if first.startswith("Query executed"):
                    continue

                def get_value(col_name: str) -> Any:
                    idx = col_map.get(col_name)
                    return row[idx] if idx is not None and idx < len(row) else None

                isin = get_value(columns["isin"]) or "UNKNOWN_ISIN"
                date_val = get_value(columns["date"])
                bo = get_value(columns["bo"]) or "UNKNOWN_BO"
                pdf_base = get_value(columns["pdf_base"])

                if not pdf_base:
                    log("WARNING", f"Skipping row: missing pdf_base in {excel_file.name}")
                    continue

                parsed_date = parse_excel_date(date_val)
                if not parsed_date:
                    log("WARNING", f"Skipping row: invalid date in {excel_file.name}")
                    continue

                record = {
                    "isin": str(isin).strip(),
                    "date": parsed_date,
                    "bo": str(bo).strip(),
                    "pdf_base": str(pdf_base).strip(),
                }

                records.append(record)
                valid_rows += 1

            log("INFO", f"{excel_file.name}: rows={file_rows}, valid={valid_rows}")

            wb.close()

        except Exception as e:
            log("ERROR", f"Failed reading Excel {excel_file}: {e}")

    log("INFO", f"Total rows processed: {total_rows}")
    return records, total_rows


def build_pdf_index(pdf_files: List[Path], case_insensitive: bool) -> Dict[str, List[Path]]:
    """Build a one-time PDF lookup index for faster record matching."""
    index: Dict[str, List[Path]] = defaultdict(list)

    for pdf_file in pdf_files:
        if not pdf_file.is_file():
            continue

        key = pdf_file.name.lower() if case_insensitive else pdf_file.name
        index[key].append(pdf_file)

    return dict(index)


def find_pdf_indexed(
    pdf_base: str, pdf_index: Dict[str, List[Path]], case_insensitive: bool
) -> List[Path]:
    """Find matching PDFs using the prebuilt index."""
    if not pdf_base.lower().endswith(".pdf"):
        pdf_base += ".pdf"

    key = pdf_base.lower() if case_insensitive else pdf_base
    return list(pdf_index.get(key, []))


# =========================
# PDF Merge
# =========================

def merge_pdfs_fast(
    pdf_paths: List[Path],
    output: Path,
    defective: List[str],
    runtime_options: Dict[str, Any],
) -> Tuple[bool, int]:
    """Try fast merge using PdfMerger."""
    try:
        merger = PdfMerger()
        appended = 0
        for p in pdf_paths:
            try:
                merger.append(str(p))
                appended += 1
            except Exception:
                defective.append(str(p))
                log("WARNING", f"Defective PDF skipped: {p}")

        if appended == 0:
            merger.close()
            return True, 0

        run_with_retries(
            f"Write merged PDF {output.name}",
            lambda: merger.write(str(output)),
            runtime_options,
        )
        merger.close()
        return True, appended
    except Exception:
        return False, 0


def merge_pdfs_fallback(
    pdf_paths: List[Path],
    output: Path,
    defective: List[str],
    runtime_options: Dict[str, Any],
) -> int:
    """Fallback merge using PdfReader/PdfWriter."""
    writer = PdfWriter()
    pages_added = 0

    for p in pdf_paths:
        try:
            reader = PdfReader(str(p))
            for page in reader.pages:
                writer.add_page(page)
                pages_added += 1
        except Exception:
            defective.append(str(p))
            log("WARNING", f"Defective PDF skipped (fallback): {p}")

    if pages_added == 0:
        return 0

    def write_output() -> None:
        with output.open("wb") as f:
            writer.write(f)

    run_with_retries(
        f"Write merged PDF {output.name}",
        write_output,
        runtime_options,
    )
    return pages_added


# =========================
# Main Processing
# =========================

def process(base_dir: Path) -> None:
    start_time = datetime.now()

    run_id = start_time.strftime("%Y-%m-%d_%H-%M-%S")
    log("INFO", f"Run started: {run_id}")

    settings = load_settings(base_dir)
    actions = settings["actions"]
    naming = settings["naming"]
    paths_cfg = settings["paths"]
    sanitize_names = settings["sanitize_filenames"]
    runtime_options = get_runtime_options(settings)
    sorting_cfg = settings.get("sorting", {})
    sorting_mode = sorting_cfg.get("mode", "default")
    bo_lookup_data = (
        load_bo_lastname_lookup(base_dir, settings, runtime_options)
        if sorting_mode == "lastname_lookup"
        else {}
    )

    log("INFO", f"Runtime mode: {runtime_options['mode']}")

    input_excel = base_dir / Path(paths_cfg["excel_input"])
    input_pdf = base_dir / Path(paths_cfg["pdf_input"])

    archive_dir = base_dir / Path(paths_cfg["archive"])
    run_dir = archive_dir / run_id

    snapshot_excel = run_dir / Path(paths_cfg["snapshot_excel"])
    snapshot_pdf = run_dir / Path(paths_cfg["snapshot_pdf"])

    individual_dir = run_dir / Path(paths_cfg["individual_output"])
    merged_dir = run_dir / Path(paths_cfg["merged_output"])

    for p in [snapshot_excel, snapshot_pdf, individual_dir, merged_dir]:
        p.mkdir(parents=True, exist_ok=True)

    excel_files = [f for f in input_excel.iterdir() if f.suffix.lower() in (".xlsx", ".xlsm")]
    pdf_files = [f for f in input_pdf.iterdir() if f.suffix.lower() == ".pdf"]

    log("INFO", f"Excel files found: {len(excel_files)}")
    log("INFO", f"PDF files found: {len(pdf_files)}")

    for f in excel_files:
        run_with_retries(
            f"Snapshot Excel copy for {f.name}",
            lambda src=f, dst=snapshot_excel / f.name: shutil.copy2(src, dst),
            runtime_options,
        )

    for f in pdf_files:
        run_with_retries(
            f"Snapshot PDF copy for {f.name}",
            lambda src=f, dst=snapshot_pdf / f.name: shutil.copy2(src, dst),
            runtime_options,
        )

    records, total_rows = read_excel_files(excel_files, settings, runtime_options)
    pdf_index = build_pdf_index(pdf_files, settings["lookup"]["case_insensitive"])

    valid_records = 0
    skipped_records = 0

    missing_pdfs: List[str] = []
    defective_pdfs: List[str] = []
    output_files: List[str] = []
    consumed_excel_files: set[Path] = set(excel_files)
    consumed_pdf_files: set[Path] = set()
    merge_source_bo_map: Dict[Path, str] = {}
    merged_all_sequence: List[Dict[str, Any]] = []

    sorted_records: List[Tuple[str, date, str, Path]] = []

    for rec in records:
        matches = find_pdf_indexed(
            rec["pdf_base"],
            pdf_index,
            settings["lookup"]["case_insensitive"],
        )

        if not matches:
            missing_pdfs.append(rec["pdf_base"])
            log("WARNING", f"Missing PDF for {rec['pdf_base']}")
            skipped_records += 1
            continue

        if len(matches) > 1:
            log("WARNING", f"Multiple PDFs found for {rec['pdf_base']}")
            skipped_records += 1
            continue

        pdf_path = matches[0]

        date_str = rec["date"].strftime(naming["date_format"])
        bo_lookup_entry = bo_lookup_data.get(normalize_lookup_key(rec["bo"]))
        sort_bo_key, display_bo = resolve_bo_sorting_values(rec["bo"], bo_lookup_entry)
        template_values = {
            "isin": str(rec["isin"]).strip(),
            "bo": str(rec["bo"]).strip(),
            "date": date_str,
            "pdf_base": str(rec["pdf_base"]).strip(),
        }

        target_dir = individual_dir
        if actions["sort_into_bo_folders"]:
            folder_name = render_template(
                naming["bo_folder_pattern"],
                {**template_values, "bo": display_bo},
                sanitize_names,
            )
            target_dir = individual_dir / folder_name

        target_dir.mkdir(parents=True, exist_ok=True)

        if actions["rename_pdfs"]:
            target_name = render_template(
                naming["output_filename_pattern"],
                template_values,
                sanitize_names,
            )
            if not target_name.lower().endswith(".pdf"):
                target_name += ".pdf"
        else:
            target_name = pdf_path.name

        target_path = ensure_unique_path(target_dir / target_name)

        try:
            source_path = snapshot_pdf / pdf_path.name
            if actions["create_outputs"]:
                run_with_retries(
                    f"Output PDF copy for {pdf_path.name}",
                    lambda src=source_path, dst=target_path: shutil.copy2(src, dst),
                    runtime_options,
                )

                log("INFO", f"Created {target_path}")
                output_files.append(str(target_path))
            else:
                log("INFO", f"Matched {pdf_path.name} without individual output creation")

            merge_source = target_path if actions["create_outputs"] else source_path
            sort_isin = template_values["isin"]

            sorted_records.append((sort_bo_key, rec["date"], sort_isin, merge_source))
            merge_source_bo_map[merge_source] = display_bo
            valid_records += 1
            consumed_pdf_files.add(pdf_path)

        except Exception as e:
            log("ERROR", f"Failed handling PDF {pdf_path}: {e}")
            skipped_records += 1

    merge_outputs: List[str] = []

    if actions["merge_all"] and sorted_records:
        sorted_records.sort(key=lambda x: (x[0].lower(), x[1], x[2].lower()))
        merged_all_sequence = [
            {
                "position": index,
                "bo": merge_source_bo_map.get(record[3], "UNKNOWN_BO"),
                "sort_key": record[0],
                "date": record[1].isoformat(),
                "isin": record[2],
                "source_pdf": str(record[3]),
            }
            for index, record in enumerate(sorted_records, start=1)
        ]

        all_paths = [r[3] for r in sorted_records]
        merged_file = merged_dir / naming["merged_all_filename"]

        fast_ok, merged_count = merge_pdfs_fast(all_paths, merged_file, defective_pdfs, runtime_options)
        if not fast_ok:
            merged_count = merge_pdfs_fallback(all_paths, merged_file, defective_pdfs, runtime_options)

        if merged_count > 0:
            merge_outputs.append(str(merged_file))
            log("INFO", f"Merged all PDFs -> {merged_file}")
        else:
            log("WARNING", "Merged all PDFs skipped: no valid PDF content available")

    if actions["per_bo_merge"]:
        per_bo: Dict[str, List[Path]] = defaultdict(list)
        for _, _, _, p in sorted_records:
            per_bo[merge_source_bo_map.get(p, "UNKNOWN_BO")].append(p)

        for bo, paths in per_bo.items():
            merged_name = render_template(
                naming["merged_bo_filename_pattern"],
                {"bo": bo},
                sanitize_names,
            )
            if not merged_name.lower().endswith(".pdf"):
                merged_name += ".pdf"

            merged_file = merged_dir / merged_name
            fast_ok, merged_count = merge_pdfs_fast(paths, merged_file, defective_pdfs, runtime_options)
            if not fast_ok:
                merged_count = merge_pdfs_fallback(paths, merged_file, defective_pdfs, runtime_options)

            if merged_count > 0:
                merge_outputs.append(str(merged_file))
                log("INFO", f"Merged BO {bo} -> {merged_file}")
            else:
                log("WARNING", f"Merged BO {bo} skipped: no valid PDF content available")

    if actions["cleanup_input"]:
        for f in consumed_excel_files:
            try:
                run_with_retries(
                    f"Delete Excel input {f.name}",
                    lambda path=f: path.unlink(),
                    runtime_options,
                )
            except Exception as e:
                log("WARNING", f"Could not delete Excel {f}: {e}")

        for f in consumed_pdf_files:
            try:
                run_with_retries(
                    f"Delete PDF input {f.name}",
                    lambda path=f: path.unlink(),
                    runtime_options,
                )
            except Exception as e:
                log("WARNING", f"Could not delete PDF {f}: {e}")

        log("INFO", "Input cleanup completed")

    unprocessed_input_pdfs = sorted(f.name for f in pdf_files if f not in consumed_pdf_files)
    if unprocessed_input_pdfs:
        log(
            "WARNING",
            "Unprocessed PDFs remain in input: " + ", ".join(unprocessed_input_pdfs),
        )

    end_time = datetime.now()
    duration_seconds = (end_time - start_time).total_seconds()
    run_status = determine_run_status(
        skipped_records,
        missing_pdfs,
        defective_pdfs,
        unprocessed_input_pdfs,
    )

    manifest = {
        "run_id": run_id,
        "status": run_status,
        "start_time": start_time.isoformat(),
        "end_time": end_time.isoformat(),
        "duration_seconds": duration_seconds,
        "excel_files": [f.name for f in excel_files],
        "pdf_input_snapshot": [f.name for f in pdf_files],
        "total_rows": total_rows,
        "valid_records": valid_records,
        "skipped_records": skipped_records,
        "missing_pdfs": missing_pdfs,
        "defective_pdfs": defective_pdfs,
        "unprocessed_input_pdf_count": len(unprocessed_input_pdfs),
        "unprocessed_input_pdfs": unprocessed_input_pdfs,
        "output_files": output_files,
        "merge_outputs": merge_outputs,
        "merged_all_sequence": merged_all_sequence,
        "settings": settings,
    }

    try:
        with (run_dir / "manifest.json").open("w", encoding="utf-8") as f:
            json.dump(manifest, f, indent=2)
    except Exception as e:
        log("ERROR", f"Failed writing manifest: {e}")

    try:
        with (run_dir / "log.txt").open("w", encoding="utf-8") as f:
            f.write("\n".join(LOG_BUFFER))
    except Exception:
        pass

    if unprocessed_input_pdfs:
        try:
            with (run_dir / "warnings.txt").open("w", encoding="utf-8") as f:
                f.write("Unprocessed PDFs remain in input:\n")
                f.write("\n".join(unprocessed_input_pdfs))
                f.write("\n")
        except Exception:
            pass

    log("INFO", f"Run finished: {run_id}")
    print_run_summary(
        run_id=run_id,
        status=run_status,
        excel_count=len(excel_files),
        pdf_count=len(pdf_files),
        total_rows=total_rows,
        valid_records=valid_records,
        skipped_records=skipped_records,
        missing_pdfs=missing_pdfs,
        defective_pdfs=defective_pdfs,
        unprocessed_input_pdfs=unprocessed_input_pdfs,
        output_files=output_files,
        merge_outputs=merge_outputs,
        run_dir=run_dir,
        duration_seconds=duration_seconds,
    )


# =========================
# Entry Point
# =========================

def main() -> None:
    """Main entry point."""
    try:
        base_dir = Path(__file__).resolve().parent
        process(base_dir)
    except Exception:
        LOG_BUFFER.append("FATAL ERROR")
        LOG_BUFFER.append(traceback.format_exc())
        try:
            fallback_log = Path("fatal_log.txt")
            with fallback_log.open("w", encoding="utf-8") as f:
                f.write("\n".join(LOG_BUFFER))
        except Exception:
            pass
        sys.exit(1)


if __name__ == "__main__":
    main()
